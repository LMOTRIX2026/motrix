import os

from flask import Flask, render_template, request, redirect, session, send_file
from flask_sqlalchemy import SQLAlchemy
from datetime import datetime, timedelta
from werkzeug.security import generate_password_hash, check_password_hash
from functools import wraps
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)

# =====================================
# CONFIGURACIÓN GENERAL
# =====================================

app.secret_key = 'motrix_secret_key'

# =====================================
# BASE DE DATOS
# =====================================

# En Render usa la variable DATABASE_URL de Neon.
# En local, si no existe DATABASE_URL, usa SQLite.
database_url = os.environ.get('DATABASE_URL', 'sqlite:///motrix.db')

if database_url.startswith('postgres://'):
    database_url = database_url.replace('postgres://', 'postgresql://', 1)

app.config['SQLALCHEMY_DATABASE_URI'] = database_url
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

db = SQLAlchemy(app)


# =====================================
# FUNCIONES DE SEGURIDAD
# =====================================

def login_requerido(funcion):
    @wraps(funcion)
    def decorada(*args, **kwargs):
        if 'usuario' not in session:
            return redirect('/login')
        return funcion(*args, **kwargs)
    return decorada


def admin_requerido(funcion):
    @wraps(funcion)
    def decorada(*args, **kwargs):
        if 'usuario' not in session:
            return redirect('/login')

        if session.get('rol') != 'admin':
            return "Acceso denegado. Esta acción solo está permitida para administradores."

        return funcion(*args, **kwargs)
    return decorada


def convertir_entero(valor):
    try:
        return int(valor)
    except:
        return 0


def convertir_fecha(valor):
    try:
        return datetime.strptime(valor, "%Y-%m-%d")
    except:
        return None


def convertir_fecha_hora(valor):
    try:
        return datetime.strptime(valor, "%Y-%m-%dT%H:%M")
    except:
        return datetime.utcnow()


def verificar_password(usuario_db, password_formulario):
    try:
        if check_password_hash(usuario_db.password, password_formulario):
            return True
    except:
        pass

    if usuario_db.password == password_formulario:
        usuario_db.password = generate_password_hash(password_formulario)
        db.session.commit()
        return True

    return False


def calcular_totales_inventario():
    productos = Producto.query.all()

    total_inventario_costo = 0
    total_inventario_venta = 0

    for producto in productos:
        stock = producto.stock or 0
        precio_compra = producto.precio_compra or 0
        precio_venta = producto.precio_venta or 0

        total_inventario_costo += stock * precio_compra
        total_inventario_venta += stock * precio_venta

    utilidad_proyectada = total_inventario_venta - total_inventario_costo

    return total_inventario_costo, total_inventario_venta, utilidad_proyectada


def calcular_porcentaje_descuento_pago(metodo_pago):
    metodo = (metodo_pago or '').lower()

    if metodo == 'datafono':
        return 5.0

    if metodo == 'transferencia':
        return 1.5

    return 0.0


def calcular_valores_venta(precio_costo, precio_venta, cantidad, metodo_pago):
    total = precio_venta * cantidad
    costo_total = precio_costo * cantidad
    ganancia_bruta = total - costo_total

    porcentaje_descuento_pago = calcular_porcentaje_descuento_pago(metodo_pago)

    # El descuento por datáfono o transferencia se calcula sobre el total pagado,
    # no sobre la ganancia.
    descuento_pago = round(total * (porcentaje_descuento_pago / 100))

    ganancia_neta = ganancia_bruta - descuento_pago

    return total, ganancia_bruta, porcentaje_descuento_pago, descuento_pago, ganancia_neta


def aplicar_filtros_ventas():
    vendedor = request.args.get('vendedor', '')
    producto = request.args.get('producto', '')
    fecha_inicio = request.args.get('fecha_inicio', '')
    fecha_fin = request.args.get('fecha_fin', '')

    consulta = Venta.query

    if vendedor:
        consulta = consulta.filter(Venta.vendedor.ilike(f"%{vendedor}%"))

    if producto:
        consulta = consulta.filter(Venta.producto.ilike(f"%{producto}%"))

    fecha_inicio_convertida = convertir_fecha(fecha_inicio)
    fecha_fin_convertida = convertir_fecha(fecha_fin)

    if fecha_inicio_convertida:
        consulta = consulta.filter(Venta.fecha >= fecha_inicio_convertida)

    if fecha_fin_convertida:
        fecha_fin_convertida = fecha_fin_convertida.replace(hour=23, minute=59, second=59)
        consulta = consulta.filter(Venta.fecha <= fecha_fin_convertida)

    ventas = consulta.order_by(Venta.fecha.desc()).all()

    return ventas, vendedor, producto, fecha_inicio, fecha_fin


def aplicar_filtros_gastos(vendedor='', fecha_inicio='', fecha_fin=''):
    consulta = Gasto.query

    if vendedor:
        consulta = consulta.filter(Gasto.registrado_por.ilike(f"%{vendedor}%"))

    fecha_inicio_convertida = convertir_fecha(fecha_inicio)
    fecha_fin_convertida = convertir_fecha(fecha_fin)

    if fecha_inicio_convertida:
        consulta = consulta.filter(Gasto.fecha >= fecha_inicio_convertida)

    if fecha_fin_convertida:
        fecha_fin_convertida = fecha_fin_convertida.replace(hour=23, minute=59, second=59)
        consulta = consulta.filter(Gasto.fecha <= fecha_fin_convertida)

    gastos = consulta.order_by(Gasto.fecha.desc()).all()

    return gastos


# =====================================
# MODELOS
# =====================================

class Usuario(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    codigo = db.Column(db.String(50))
    nombre = db.Column(db.String(100))
    usuario = db.Column(db.String(100), unique=True)
    password = db.Column(db.String(255))
    rol = db.Column(db.String(50))


class Producto(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    referencia = db.Column(db.String(100))
    nombre = db.Column(db.String(200))
    marca = db.Column(db.String(100))
    categoria = db.Column(db.String(100))
    stock = db.Column(db.Integer)
    precio_compra = db.Column(db.Integer)
    precio_venta = db.Column(db.Integer)


class Venta(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    producto_id = db.Column(db.Integer)
    producto = db.Column(db.String(200))
    cantidad = db.Column(db.Integer)
    precio_costo = db.Column(db.Integer)
    precio_venta = db.Column(db.Integer)
    total = db.Column(db.Integer)
    ganancia_bruta = db.Column(db.Integer)
    metodo_pago = db.Column(db.String(50))
    porcentaje_descuento_pago = db.Column(db.Float)
    descuento_pago = db.Column(db.Integer)
    ganancia_neta = db.Column(db.Integer)
    vendedor = db.Column(db.String(100))
    fecha = db.Column(db.DateTime, default=datetime.utcnow)


class Gasto(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    descripcion = db.Column(db.String(250))
    valor = db.Column(db.Integer)
    registrado_por = db.Column(db.String(100))
    fecha = db.Column(db.DateTime, default=datetime.utcnow)


# =====================================
# LOGIN / LOGOUT
# =====================================

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        usuario = request.form['usuario']
        password = request.form['password']

        user = Usuario.query.filter_by(usuario=usuario).first()

        if user and verificar_password(user, password):
            session['usuario'] = user.nombre
            session['rol'] = user.rol
            session['usuario_id'] = user.id
            return redirect('/')

        return "Usuario o contraseña incorrectos"

    return render_template('login.html')


@app.route('/logout')
def logout():
    session.clear()
    return redirect('/login')


# =====================================
# DASHBOARD
# =====================================

@app.route('/')
@login_requerido
def inicio():
    productos = Producto.query.count()

    if session.get('rol') == 'admin':
        ventas = Venta.query.count()
        total_gastos = db.session.query(db.func.sum(Gasto.valor)).scalar() or 0
    else:
        ventas = Venta.query.filter_by(vendedor=session.get('usuario')).count()
        total_gastos = db.session.query(db.func.sum(Gasto.valor)).filter_by(registrado_por=session.get('usuario')).scalar() or 0

    total_ventas = db.session.query(db.func.sum(Venta.total)).scalar() or 0
    total_ganancias = db.session.query(db.func.sum(Venta.ganancia_neta)).scalar() or 0

    total_inventario_costo, total_inventario_venta, utilidad_proyectada = calcular_totales_inventario()

    return render_template(
        'inicio.html',
        productos=productos,
        ventas=ventas,
        total_ventas=total_ventas,
        total_ganancias=total_ganancias,
        total_gastos=total_gastos,
        total_inventario_costo=total_inventario_costo,
        total_inventario_venta=total_inventario_venta,
        utilidad_proyectada=utilidad_proyectada
    )


# =====================================
# USUARIOS
# =====================================

@app.route('/usuarios')
@admin_requerido
def usuarios():
    usuarios = Usuario.query.all()
    return render_template('usuarios.html', usuarios=usuarios)


@app.route('/agregar_usuario', methods=['POST'])
@admin_requerido
def agregar_usuario():
    usuario_existente = Usuario.query.filter_by(usuario=request.form['usuario']).first()

    if usuario_existente:
        return "El nombre de usuario ya existe"

    nuevo_usuario = Usuario(
        codigo=request.form['codigo'],
        nombre=request.form['nombre'],
        usuario=request.form['usuario'],
        password=generate_password_hash(request.form['password']),
        rol=request.form['rol']
    )

    db.session.add(nuevo_usuario)
    db.session.commit()

    return redirect('/usuarios')


@app.route('/eliminar_usuario/<int:id>')
@admin_requerido
def eliminar_usuario(id):
    usuario = Usuario.query.get_or_404(id)

    if session.get('usuario_id') == usuario.id:
        return "No puedes eliminar tu propio usuario mientras tienes la sesión iniciada"

    db.session.delete(usuario)
    db.session.commit()

    return redirect('/usuarios')


# =====================================
# INVENTARIO
# =====================================

@app.route('/inventario')
@login_requerido
def inventario():
    productos = Producto.query.order_by(Producto.id.desc()).all()
    total_inventario_costo, total_inventario_venta, utilidad_proyectada = calcular_totales_inventario()

    return render_template(
        'inventario.html',
        productos=productos,
        total_inventario_costo=total_inventario_costo,
        total_inventario_venta=total_inventario_venta,
        utilidad_proyectada=utilidad_proyectada
    )


@app.route('/agregar_producto', methods=['POST'])
@admin_requerido
def agregar_producto():
    nuevo_producto = Producto(
        referencia=request.form['referencia'],
        nombre=request.form['nombre'],
        marca=request.form['marca'],
        categoria=request.form['categoria'],
        stock=convertir_entero(request.form['stock']),
        precio_compra=convertir_entero(request.form['precio_compra']),
        precio_venta=convertir_entero(request.form['precio_venta'])
    )

    db.session.add(nuevo_producto)
    db.session.commit()

    return redirect('/inventario')


@app.route('/eliminar_producto/<int:id>')
@admin_requerido
def eliminar_producto(id):
    producto = Producto.query.get_or_404(id)
    db.session.delete(producto)
    db.session.commit()
    return redirect('/inventario')


@app.route('/editar_producto/<int:id>', methods=['GET', 'POST'])
@admin_requerido
def editar_producto(id):
    producto = Producto.query.get_or_404(id)

    if request.method == 'POST':
        producto.referencia = request.form['referencia']
        producto.nombre = request.form['nombre']
        producto.marca = request.form['marca']
        producto.categoria = request.form['categoria']
        producto.stock = convertir_entero(request.form['stock'])
        producto.precio_compra = convertir_entero(request.form['precio_compra'])
        producto.precio_venta = convertir_entero(request.form['precio_venta'])

        db.session.commit()

        return redirect('/inventario')

    return render_template('editar_producto.html', producto=producto)


# =====================================
# VENTAS
# =====================================

@app.route('/ventas')
@login_requerido
def ventas():
    productos = Producto.query.order_by(Producto.nombre.asc()).all()

    if session.get('rol') == 'admin':
        ventas = Venta.query.order_by(Venta.id.desc()).all()
    else:
        ventas = Venta.query.filter_by(vendedor=session.get('usuario')).order_by(Venta.id.desc()).all()

    return render_template('ventas.html', productos=productos, ventas=ventas)


@app.route('/registrar_venta', methods=['POST'])
@login_requerido
def registrar_venta():
    producto_id = convertir_entero(request.form['producto_id'])
    cantidad = convertir_entero(request.form['cantidad'])
    precio_venta_manual = convertir_entero(request.form['precio_venta_manual'])
    metodo_pago = request.form['metodo_pago']
    fecha_venta = convertir_fecha_hora(request.form.get('fecha_venta', ''))

    if cantidad <= 0:
        return "La cantidad debe ser mayor que cero"

    if precio_venta_manual <= 0:
        return "Debes ingresar un precio de venta válido"

    if metodo_pago not in ['efectivo', 'transferencia', 'datafono']:
        return "Debes seleccionar un método de pago válido"

    producto = Producto.query.get(producto_id)

    if not producto:
        return "Producto no encontrado"

    if cantidad > producto.stock:
        return "No hay suficiente stock"

    producto.stock -= cantidad

    total, ganancia_bruta, porcentaje_descuento_pago, descuento_pago, ganancia_neta = calcular_valores_venta(
        producto.precio_compra,
        precio_venta_manual,
        cantidad,
        metodo_pago
    )

    nueva_venta = Venta(
        producto_id=producto.id,
        producto=producto.nombre,
        cantidad=cantidad,
        precio_costo=producto.precio_compra,
        precio_venta=precio_venta_manual,
        total=total,
        ganancia_bruta=ganancia_bruta,
        metodo_pago=metodo_pago,
        porcentaje_descuento_pago=porcentaje_descuento_pago,
        descuento_pago=descuento_pago,
        ganancia_neta=ganancia_neta,
        vendedor=session['usuario'],
        fecha=fecha_venta
    )

    db.session.add(nueva_venta)
    db.session.commit()

    return redirect('/ventas')


@app.route('/eliminar_venta/<int:id>')
@admin_requerido
def eliminar_venta(id):
    venta = Venta.query.get_or_404(id)
    producto = Producto.query.get(venta.producto_id)

    if producto:
        producto.stock += venta.cantidad

    db.session.delete(venta)
    db.session.commit()

    return redirect('/ventas')


@app.route('/editar_venta/<int:id>', methods=['GET', 'POST'])
@admin_requerido
def editar_venta(id):
    venta = Venta.query.get_or_404(id)
    producto = Producto.query.get(venta.producto_id)

    if request.method == 'POST':
        nueva_cantidad = convertir_entero(request.form['cantidad'])
        nuevo_precio_venta = convertir_entero(request.form['precio_venta'])
        nuevo_metodo_pago = request.form['metodo_pago']
        nueva_fecha = convertir_fecha_hora(request.form.get('fecha_venta', ''))

        if nueva_cantidad <= 0:
            return "La cantidad debe ser mayor que cero"

        if nuevo_precio_venta <= 0:
            return "El precio de venta debe ser mayor que cero"

        if nuevo_metodo_pago not in ['efectivo', 'transferencia', 'datafono']:
            return "Debes seleccionar un método de pago válido"

        if producto:
            diferencia = nueva_cantidad - venta.cantidad

            if diferencia > producto.stock:
                return "No hay suficiente stock para aumentar esta venta"

            producto.stock -= diferencia

        total, ganancia_bruta, porcentaje_descuento_pago, descuento_pago, ganancia_neta = calcular_valores_venta(
            venta.precio_costo,
            nuevo_precio_venta,
            nueva_cantidad,
            nuevo_metodo_pago
        )

        venta.cantidad = nueva_cantidad
        venta.precio_venta = nuevo_precio_venta
        venta.total = total
        venta.ganancia_bruta = ganancia_bruta
        venta.metodo_pago = nuevo_metodo_pago
        venta.porcentaje_descuento_pago = porcentaje_descuento_pago
        venta.descuento_pago = descuento_pago
        venta.ganancia_neta = ganancia_neta
        venta.fecha = nueva_fecha

        db.session.commit()

        return redirect('/ventas')

    return render_template('editar_venta.html', venta=venta)


# =====================================
# GASTOS
# =====================================

@app.route('/gastos')
@login_requerido
def gastos():
    if session.get('rol') == 'admin':
        gastos = Gasto.query.order_by(Gasto.fecha.desc()).all()
    else:
        gastos = Gasto.query.filter_by(registrado_por=session.get('usuario')).order_by(Gasto.fecha.desc()).all()

    total_gastos = sum((gasto.valor or 0) for gasto in gastos)

    return render_template(
        'gastos.html',
        gastos=gastos,
        total_gastos=total_gastos
    )


@app.route('/registrar_gasto', methods=['POST'])
@login_requerido
def registrar_gasto():
    descripcion = request.form.get('descripcion', '').strip()
    valor = convertir_entero(request.form.get('valor', '0'))
    fecha_gasto = convertir_fecha_hora(request.form.get('fecha_gasto', ''))

    if not descripcion:
        return "Debes ingresar la descripción del gasto"

    if valor <= 0:
        return "Debes ingresar un valor de gasto válido"

    nuevo_gasto = Gasto(
        descripcion=descripcion,
        valor=valor,
        registrado_por=session.get('usuario'),
        fecha=fecha_gasto
    )

    db.session.add(nuevo_gasto)
    db.session.commit()

    return redirect('/gastos')


@app.route('/eliminar_gasto/<int:id>')
@admin_requerido
def eliminar_gasto(id):
    gasto = Gasto.query.get_or_404(id)
    db.session.delete(gasto)
    db.session.commit()
    return redirect('/gastos')


# =====================================
# INFORMES
# =====================================

@app.route('/informes')
@admin_requerido
def informes():
    ventas, vendedor, producto, fecha_inicio, fecha_fin = aplicar_filtros_ventas()

    gastos = aplicar_filtros_gastos(vendedor, fecha_inicio, fecha_fin)

    costo_fijo_nombre = request.args.get('costo_fijo_nombre', '')
    costo_fijo_valor = convertir_entero(request.args.get('costo_fijo_valor', '0'))

    dias_proyeccion = convertir_entero(request.args.get('dias_proyeccion', '30'))

    total_vendido = sum((venta.total or 0) for venta in ventas)
    total_ganancia_bruta = sum((venta.ganancia_bruta or 0) for venta in ventas)
    total_descuento_pago = sum((venta.descuento_pago or 0) for venta in ventas)
    total_ganancia_neta_antes_costos = sum((venta.ganancia_neta or 0) for venta in ventas)
    total_cantidad = sum((venta.cantidad or 0) for venta in ventas)

    total_gastos_dia = sum((gasto.valor or 0) for gasto in gastos)

    ganancia_final = total_ganancia_neta_antes_costos - costo_fijo_valor - total_gastos_dia

    fechas_unicas = set()
    for venta in ventas:
        if venta.fecha:
            fechas_unicas.add(venta.fecha.date())

    dias_con_ventas = len(fechas_unicas)

    if dias_con_ventas > 0:
        promedio_dia_venta = round(total_vendido / dias_con_ventas)
        promedio_dia_ganancia = round(ganancia_final / dias_con_ventas)
    else:
        promedio_dia_venta = 0
        promedio_dia_ganancia = 0

    proyeccion_ventas = promedio_dia_venta * dias_proyeccion
    proyeccion_ganancia = promedio_dia_ganancia * dias_proyeccion

    vendedores = db.session.query(Venta.vendedor).distinct().all()
    productos = db.session.query(Venta.producto).distinct().all()

    return render_template(
        'informes.html',
        ventas=ventas,
        gastos=gastos,
        total_vendido=total_vendido,
        total_ganancia_bruta=total_ganancia_bruta,
        total_descuento_pago=total_descuento_pago,
        total_ganancia_neta_antes_costos=total_ganancia_neta_antes_costos,
        total_cantidad=total_cantidad,
        total_gastos_dia=total_gastos_dia,
        costo_fijo_nombre=costo_fijo_nombre,
        costo_fijo_valor=costo_fijo_valor,
        ganancia_final=ganancia_final,
        dias_con_ventas=dias_con_ventas,
        promedio_dia_venta=promedio_dia_venta,
        promedio_dia_ganancia=promedio_dia_ganancia,
        dias_proyeccion=dias_proyeccion,
        proyeccion_ventas=proyeccion_ventas,
        proyeccion_ganancia=proyeccion_ganancia,
        vendedores=vendedores,
        productos=productos,
        vendedor=vendedor,
        producto=producto,
        fecha_inicio=fecha_inicio,
        fecha_fin=fecha_fin
    )


@app.route('/descargar_informe_excel')
@admin_requerido
def descargar_informe_excel():
    ventas, vendedor, producto, fecha_inicio, fecha_fin = aplicar_filtros_ventas()
    gastos = aplicar_filtros_gastos(vendedor, fecha_inicio, fecha_fin)

    costo_fijo_nombre = request.args.get('costo_fijo_nombre', '')
    costo_fijo_valor = convertir_entero(request.args.get('costo_fijo_valor', '0'))
    dias_proyeccion = convertir_entero(request.args.get('dias_proyeccion', '30'))

    total_vendido = sum((v.total or 0) for v in ventas)
    total_ganancia_bruta = sum((v.ganancia_bruta or 0) for v in ventas)
    total_descuento_pago = sum((v.descuento_pago or 0) for v in ventas)
    total_ganancia_neta_antes_costos = sum((v.ganancia_neta or 0) for v in ventas)
    total_cantidad = sum((v.cantidad or 0) for v in ventas)
    total_gastos_dia = sum((g.valor or 0) for g in gastos)
    ganancia_final = total_ganancia_neta_antes_costos - costo_fijo_valor - total_gastos_dia

    fechas_unicas = set()
    for venta in ventas:
        if venta.fecha:
            fechas_unicas.add(venta.fecha.date())

    dias_con_ventas = len(fechas_unicas)

    if dias_con_ventas > 0:
        promedio_dia_venta = round(total_vendido / dias_con_ventas)
        promedio_dia_ganancia = round(ganancia_final / dias_con_ventas)
    else:
        promedio_dia_venta = 0
        promedio_dia_ganancia = 0

    proyeccion_ventas = promedio_dia_venta * dias_proyeccion
    proyeccion_ganancia = promedio_dia_ganancia * dias_proyeccion

    archivo = Workbook()
    hoja = archivo.active
    hoja.title = "Informe de ventas"

    hoja.append(["INFORME DE VENTAS MOTRIX"])
    hoja.append([])
    hoja.append(["Filtro vendedor", vendedor or "Todos"])
    hoja.append(["Filtro producto", producto or "Todos"])
    hoja.append(["Fecha inicial", fecha_inicio or "Sin filtro"])
    hoja.append(["Fecha final", fecha_fin or "Sin filtro"])
    hoja.append(["Costo fijo", costo_fijo_nombre or "No registrado"])
    hoja.append(["Valor costo fijo", costo_fijo_valor])
    hoja.append(["Gastos registrados", total_gastos_dia])
    hoja.append(["Días con ventas", dias_con_ventas])
    hoja.append(["Promedio diario ventas", promedio_dia_venta])
    hoja.append(["Promedio diario ganancia", promedio_dia_ganancia])
    hoja.append(["Días de proyección", dias_proyeccion])
    hoja.append(["Proyección ventas", proyeccion_ventas])
    hoja.append(["Proyección ganancia", proyeccion_ganancia])
    hoja.append([])

    encabezados = [
        "ID", "Fecha", "Producto", "Cantidad", "Precio costo", "Precio venta",
        "Total vendido", "Ganancia bruta", "Método de pago", "% descuento pago",
        "Descuento pago", "Ganancia neta", "Vendedor"
    ]

    hoja.append(encabezados)

    for venta in ventas:
        hoja.append([
            venta.id,
            venta.fecha.strftime('%d/%m/%Y %H:%M'),
            venta.producto,
            venta.cantidad,
            venta.precio_costo,
            venta.precio_venta,
            venta.total,
            venta.ganancia_bruta,
            venta.metodo_pago,
            venta.porcentaje_descuento_pago,
            venta.descuento_pago,
            venta.ganancia_neta,
            venta.vendedor
        ])

    hoja.append([])
    hoja.append(["GASTOS REGISTRADOS"])
    hoja.append(["ID", "Fecha", "Descripción", "Valor", "Registrado por"])

    for gasto in gastos:
        hoja.append([
            gasto.id,
            gasto.fecha.strftime('%d/%m/%Y %H:%M'),
            gasto.descripcion,
            gasto.valor,
            gasto.registrado_por
        ])

    hoja.append([])
    hoja.append(["RESUMEN"])
    hoja.append(["TOTAL CANTIDAD", total_cantidad])
    hoja.append(["TOTAL VENDIDO", total_vendido])
    hoja.append(["GANANCIA BRUTA", total_ganancia_bruta])
    hoja.append(["DESCUENTOS POR MÉTODO DE PAGO", total_descuento_pago])
    hoja.append(["GANANCIA NETA ANTES DE COSTOS Y GASTOS", total_ganancia_neta_antes_costos])
    hoja.append(["COSTO FIJO", costo_fijo_valor])
    hoja.append(["GASTOS REGISTRADOS", total_gastos_dia])
    hoja.append(["GANANCIA FINAL", ganancia_final])
    hoja.append(["PROMEDIO DIARIO DE VENTAS", promedio_dia_venta])
    hoja.append(["PROMEDIO DIARIO DE GANANCIA", promedio_dia_ganancia])
    hoja.append(["PROYECCIÓN DE VENTAS", proyeccion_ventas])
    hoja.append(["PROYECCIÓN DE GANANCIA", proyeccion_ganancia])

    color_encabezado = PatternFill(start_color="111827", end_color="111827", fill_type="solid")
    fuente_encabezado = Font(color="FFFFFF", bold=True)
    borde = Border(
        left=Side(style='thin', color='D1D5DB'),
        right=Side(style='thin', color='D1D5DB'),
        top=Side(style='thin', color='D1D5DB'),
        bottom=Side(style='thin', color='D1D5DB')
    )

    for fila in hoja.iter_rows():
        for celda in fila:
            celda.border = borde
            celda.alignment = Alignment(vertical='center')

    hoja[1][0].font = Font(bold=True, size=16)

    for celda in hoja[17]:
        celda.fill = color_encabezado
        celda.font = fuente_encabezado
        celda.alignment = Alignment(horizontal='center')

    for columna in hoja.columns:
        maximo = 0
        letra = get_column_letter(columna[0].column)

        for celda in columna:
            if celda.value:
                maximo = max(maximo, len(str(celda.value)))

        hoja.column_dimensions[letra].width = maximo + 4

    memoria = BytesIO()
    archivo.save(memoria)
    memoria.seek(0)

    return send_file(
        memoria,
        as_attachment=True,
        download_name="informe_ventas_motrix.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


# =====================================
# CREAR USUARIO ADMIN AUTOMÁTICO
# =====================================

def crear_admin():
    admin = Usuario.query.filter_by(usuario='admin').first()

    if not admin:
        nuevo_admin = Usuario(
            codigo='ADM001',
            nombre='Administrador',
            usuario='admin',
            password=generate_password_hash('Admin2026'),
            rol='admin'
        )

        db.session.add(nuevo_admin)
        db.session.commit()


# =====================================
# CREAR BASE DE DATOS AL INICIAR
# =====================================

with app.app_context():
    db.create_all()
    crear_admin()


# =====================================
# EJECUTAR LOCALMENTE
# =====================================

if __name__ == '__main__':
    app.run(debug=True)
